import os
import sqlite3
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_PATH = r"d:\SM0237 Onedrive\OneDrive - TVS Mobility Private Limited\Documents\PROJECT\data\overall_catalogue.db"
OUTPUT_CSV = r"d:\SM0237 Onedrive\OneDrive - TVS Mobility Private Limited\Documents\PROJECT\Exact_100pct_Duplicate_Rows_All_21_Columns.csv"
OUTPUT_EXCEL = r"d:\SM0237 Onedrive\OneDrive - TVS Mobility Private Limited\Documents\PROJECT\Executive_Exact_Duplicate_Catalog_Report.xlsx"

print("1. Classification & Sheet 1 Generation (Brand-Wise Breakdown with OEM/OES)...", flush=True)

conn = sqlite3.connect(DB_PATH)

# OEM Brands list
oem_brands = {
    'HYUNDAI', 'MARUTI SUZUKI', 'TATA', 'MAHINDRA', 'HONDA', 'TOYOTA', 'FORD', 
    'VOLKSWAGEN', 'SKODA', 'CHEVROLET', 'NISSAN', 'RENAULT', 'AUDI', 'BMW', 
    'MERCEDES BENZ', 'KIA', 'JEEP', 'FIAT', 'FORCE', 'VOLVO', 'NEXA', 
    'MORRIS GARAGES', 'TATA FICOSA', 'MSL - MAHINDRA SONA', 'SONA MAHINDRA'
}

# Consumables / Oils & Paints
consumable_brands = {
    'ASIAN PPG TINTS', 'CASTROL', 'SHELL LUBES', '3M', 'WURTH', 'ANABOND', 
    'KWIK PATCH', 'PAINTS AND CONSUMABLES', 'ADD LUB ITEMS', 'REFEX'
}

query_sheet1 = """
    SELECT 
        brand_name,
        COUNT(DISTINCT line_code_part_number) AS unique_sku_count,
        COUNT(*) AS total_db_records
    FROM overall_catalogue
    WHERE brand_name IS NOT NULL AND brand_name != ''
    GROUP BY brand_name
    ORDER BY total_db_records DESC
"""
df_s1 = pd.read_sql_query(query_sheet1, conn)

total_db = df_s1['total_db_records'].sum()

def categorize_brand(brand):
    b_upper = str(brand).upper().strip()
    if b_upper in oem_brands:
        return 'OEM (Vehicle Manufacturer)'
    elif b_upper in consumable_brands:
        return 'Consumables & Lubricants'
    elif any(kw in b_upper for kw in ['LOCAL', 'MYTVS', 'ITEMS', 'MONOGRAM', 'RUBBER', 'GENERAL']):
        return 'Aftermarket & Local Spares'
    else:
        return 'OES (Equipment Supplier)'

df_s1['Category (OEM/OES)'] = df_s1['brand_name'].apply(categorize_brand)
df_s1['Contribution %'] = (df_s1['total_db_records'] / total_db) * 100

df_s1 = df_s1[['brand_name', 'Category (OEM/OES)', 'unique_sku_count', 'total_db_records', 'Contribution %']]
df_s1.columns = ['Brand Name', 'Brand Category (OEM/OES)', 'Unique SKU Count', 'Total DB Records', 'Contribution %']

print(f"Sheet 1 compiled: {len(df_s1)} brands, Total DB Records = {total_db:,}, Total Unique SKUs = {df_s1['Unique SKU Count'].sum():,}", flush=True)

print("\n2. Compiling Sheet 2 & CSV: All 21 Columns Exact Duplicate Rows + Repeated Count...", flush=True)

query_sheet2 = """
    SELECT 
        line_code,
        brand_name,
        line_code_part_number,
        part_description,
        hsn_code,
        tax_rate,
        price_type,
        sales_uom,
        list_price,
        mrp_price,
        sales_price_group,
        vehicle_segment,
        vehicle_make,
        vehicle_model,
        vehicle_sub_model,
        vehicle_fuel_type,
        vehicle_from_year,
        vehicle_to_year,
        aggregate,
        sub_aggregate,
        components,
        COUNT(*) AS Repeated_Count_In_DB
    FROM overall_catalogue
    WHERE brand_name IS NOT NULL AND brand_name != ''
    GROUP BY 
        line_code, brand_name, line_code_part_number, part_description,
        hsn_code, tax_rate, price_type, sales_uom, list_price, mrp_price,
        sales_price_group, vehicle_segment, vehicle_make, vehicle_model,
        vehicle_sub_model, vehicle_fuel_type, vehicle_from_year, vehicle_to_year,
        aggregate, sub_aggregate, components
    HAVING Repeated_Count_In_DB > 1
    ORDER BY Repeated_Count_In_DB DESC
"""

df_s2 = pd.read_sql_query(query_sheet2, conn)
conn.close()

print(f"Found {len(df_s2):,} unique duplicate row patterns.", flush=True)

print(f"\n3. Exporting full {len(df_s2):,} duplicate rows to CSV...", flush=True)
df_s2.to_csv(OUTPUT_CSV, index=False)
csv_size_mb = round(os.path.getsize(OUTPUT_CSV)/(1024*1024), 2)
print(f"CSV Export Complete! ({csv_size_mb} MB)", flush=True)

df_s2_excel = df_s2.head(500000)
print(f"\n4. Exporting Top {len(df_s2_excel):,} duplicate rows to Excel with styling...", flush=True)

navy_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
red_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
red_bold = Font(name="Calibri", size=11, bold=True, color="C00000")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

with pd.ExcelWriter(OUTPUT_EXCEL, engine='openpyxl') as writer:
    df_s1.to_excel(writer, sheet_name='Brand Master & OEM Summary', index=False)
    df_s2_excel.to_excel(writer, sheet_name='Top Duplicate Rows Report', index=False)

    ws1 = writer.sheets['Brand Master & OEM Summary']
    ws2 = writer.sheets['Top Duplicate Rows Report']

    ws1.views.sheetView[0].showGridLines = True
    ws2.views.sheetView[0].showGridLines = True

    # Format Sheet 1
    for col in range(1, ws1.max_column + 1):
        cell = ws1.cell(row=1, column=col)
        cell.fill = navy_fill
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in range(2, ws1.max_row + 1):
        ws1.cell(row=row, column=3).number_format = '#,##0'
        ws1.cell(row=row, column=4).number_format = '#,##0'
        ws1.cell(row=row, column=5).number_format = '0.00"%"'
        for col in range(1, ws1.max_column + 1):
            ws1.cell(row=row, column=col).border = thin_border

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 16)

    # Format Sheet 2 Headers
    for col in range(1, ws2.max_column + 1):
        cell = ws2.cell(row=1, column=col)
        cell.fill = red_fill if col == ws2.max_column else navy_fill
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Format Sheet 2 Column 22 number_format
    rep_col_idx = ws2.max_column
    for row in range(2, ws2.max_row + 1):
        c = ws2.cell(row=row, column=rep_col_idx)
        c.number_format = '#,##0'
        c.font = red_bold

    # Column widths for Sheet 2
    for col_idx in range(1, ws2.max_column + 1):
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = 22

excel_size_mb = round(os.path.getsize(OUTPUT_EXCEL)/(1024*1024), 2)

print("\n========================================================", flush=True)
print("SUCCESS! Reports Regenerated with Clean 167 Brands!", flush=True)
print(f"Excel Report: {OUTPUT_EXCEL} ({excel_size_mb} MB)", flush=True)
print(f"CSV Complete Report: {OUTPUT_CSV} ({csv_size_mb} MB)", flush=True)
print("========================================================", flush=True)
