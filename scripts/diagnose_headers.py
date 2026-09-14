import os
import glob
import openpyxl

stock_folder = r"ALL GOOD STOCK"
files = sorted(glob.glob(os.path.join(stock_folder, "*.xlsx")) + glob.glob(os.path.join(stock_folder, "*.csv")))

print(f"Diagnosing headers for {len(files)} files in ALL GOOD STOCK...\n")

for fpath in files:
    fname = os.path.basename(fpath)
    wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
    ws = wb.active
    
    headers = []
    for row in ws.iter_rows(values_only=True):
        headers = [str(c or '').strip() for c in row]
        break
    wb.close()

    print(f"=== {fname} ({len(headers)} columns) ===")
    for idx, h in enumerate(headers):
        if h:
            print(f"  Col {idx} (Excel {openpyxl.utils.get_column_letter(idx+1)}): '{h}'")
    print("\n")
